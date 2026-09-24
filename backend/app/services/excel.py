"""Exportación del reporte diario de ventas a Excel (.xlsx).

El archivo se entrega con las columnas separadas, tipos numéricos reales
(no texto) y un autofiltro activo, para que quien lo reciba pueda ordenar,
filtrar y hacer tablas dinámicas sin limpiar nada antes.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import get_settings

settings = get_settings()

VERDE = "FF2B3A2F"
ORO = "FFC9A227"
PAPEL = "FFF7F4EC"
FRANJA = "FFEFEADE"
FORMATO_MONEDA = '"$" #,##0.00'

_borde_fino = Border(bottom=Side(style="thin", color="FFD8D2C4"))

CABECERAS = [
    ("N.º venta", 18), ("Fecha", 12), ("Hora", 9), ("Cliente", 26),
    ("Documento", 16), ("Productos / servicios", 46), ("Unidades", 11),
    ("Subtotal", 15), ("Descuento", 14), ("IVA", 14), ("Total", 16),
    ("Método de pago", 16), ("Estado", 16), ("Factura", 14),
]


def generar_reporte_ventas_excel(ventas: list, dia) -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Ventas"

    # ── Encabezado del documento ─────────────────────────────────────────
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CABECERAS))
    titulo = hoja.cell(row=1, column=1, value=f"{settings.empresa_nombre} · Reporte diario de ventas")
    titulo.font = Font(bold=True, size=14, color=PAPEL)
    titulo.fill = PatternFill("solid", fgColor=VERDE)
    titulo.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    hoja.row_dimensions[1].height = 26

    hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(CABECERAS))
    subtitulo = hoja.cell(
        row=2, column=1,
        value=f"Fecha del reporte: {dia.strftime('%d/%m/%Y')}  ·  "
              f"NIT {settings.empresa_nit}  ·  {settings.empresa_email}",
    )
    subtitulo.font = Font(size=9, color="FF6B675F")
    subtitulo.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # ── Fila de cabeceras ────────────────────────────────────────────────
    fila_cabecera = 4
    for indice, (texto, ancho) in enumerate(CABECERAS, start=1):
        celda = hoja.cell(row=fila_cabecera, column=indice, value=texto)
        celda.font = Font(bold=True, size=9, color=PAPEL)
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.row_dimensions[fila_cabecera].height = 24

    # ── Datos ────────────────────────────────────────────────────────────
    fila = fila_cabecera + 1
    for indice, v in enumerate(ventas):
        items = ", ".join(d.descripcion for d in v.detalles)
        cliente = f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "—"
        documento = (
            f"{v.cliente.tipo_documento} {v.cliente.numero_documento}"
            if v.cliente and getattr(v.cliente, "numero_documento", None)
            else ""
        )
        valores = [
            v.numero,
            v.creado_en.strftime("%d/%m/%Y"),
            v.creado_en.strftime("%H:%M"),
            cliente,
            documento,
            items,
            sum(d.cantidad for d in v.detalles),
            float(v.subtotal),
            float(v.descuento),
            float(v.impuestos),
            float(v.total),
            v.metodo_pago.value.capitalize(),
            v.estado.value.replace("_", " ").capitalize(),
            v.factura.numero if getattr(v, "factura", None) else "",
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)
            celda.font = Font(size=9)
            celda.border = _borde_fino
            if indice % 2 == 0:
                celda.fill = PatternFill("solid", fgColor=FRANJA)
            if columna in (8, 9, 10, 11):
                celda.number_format = FORMATO_MONEDA
                celda.alignment = Alignment(horizontal="right")
            elif columna == 7:
                celda.alignment = Alignment(horizontal="center")
        fila += 1

    if not ventas:
        hoja.cell(row=fila, column=1, value=f"No se registraron ventas el {dia.strftime('%d/%m/%Y')}.")
        hoja.cell(row=fila, column=1).font = Font(italic=True, size=9, color="FF6B675F")
        fila += 1

    # ── Fila de totales ──────────────────────────────────────────────────
    fila_total = fila
    etiqueta = hoja.cell(row=fila_total, column=6, value="TOTAL DEL DÍA")
    etiqueta.font = Font(bold=True, size=10, color=PAPEL)
    etiqueta.fill = PatternFill("solid", fgColor=VERDE)
    etiqueta.alignment = Alignment(horizontal="right")

    for columna in (7, 8, 9, 10, 11):
        letra = get_column_letter(columna)
        primera, ultima = fila_cabecera + 1, fila_total - 1
        # Fórmula real: quien abra el archivo ve la suma, no un valor pegado.
        formula = f"=SUM({letra}{primera}:{letra}{ultima})" if ventas else 0
        celda = hoja.cell(row=fila_total, column=columna, value=formula)
        celda.font = Font(bold=True, size=10, color=PAPEL)
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="right" if columna != 7 else "center")
        if columna != 7:
            celda.number_format = FORMATO_MONEDA

    hoja.cell(row=fila_total, column=1).fill = PatternFill("solid", fgColor=VERDE)
    for columna in range(1, 6):
        hoja.cell(row=fila_total, column=columna).fill = PatternFill("solid", fgColor=VERDE)
    for columna in range(12, len(CABECERAS) + 1):
        hoja.cell(row=fila_total, column=columna).fill = PatternFill("solid", fgColor=VERDE)

    # ── Usabilidad: filtros y paneles congelados ─────────────────────────
    if ventas:
        hoja.auto_filter.ref = (
            f"A{fila_cabecera}:{get_column_letter(len(CABECERAS))}{fila_total - 1}"
        )
    hoja.freeze_panes = f"A{fila_cabecera + 1}"

    # ── Segunda hoja: resumen por estado ─────────────────────────────────
    resumen = wb.create_sheet("Resumen")
    resumen.column_dimensions["A"].width = 24
    resumen.column_dimensions["B"].width = 14
    resumen.column_dimensions["C"].width = 18

    for columna, texto in enumerate(["Estado", "Ventas", "Total"], start=1):
        celda = resumen.cell(row=1, column=columna, value=texto)
        celda.font = Font(bold=True, size=10, color=PAPEL)
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="center")

    por_estado: dict[str, list[float]] = {}
    for v in ventas:
        clave = v.estado.value.replace("_", " ").capitalize()
        acumulado = por_estado.setdefault(clave, [0, 0.0])
        acumulado[0] += 1
        acumulado[1] += float(v.total)

    fila_resumen = 2
    for estado, (conteo, monto) in sorted(por_estado.items()):
        resumen.cell(row=fila_resumen, column=1, value=estado).font = Font(size=9)
        resumen.cell(row=fila_resumen, column=2, value=conteo).alignment = Alignment(horizontal="center")
        celda_total = resumen.cell(row=fila_resumen, column=3, value=monto)
        celda_total.number_format = FORMATO_MONEDA
        fila_resumen += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
